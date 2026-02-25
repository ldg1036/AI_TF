import argparse
import json
import os
from datetime import datetime

from openpyxl import load_workbook

VALID_STATUS = {"OK", "NG", "N/A", "SKIP"}


def _detect_status_column(ws):
    for row in range(1, min(ws.max_row, 50) + 1):
        for col in range(1, min(ws.max_column, 20) + 1):
            value = str(ws.cell(row=row, column=col).value or "").strip()
            if not value:
                continue
            if "3차 검증" in value:
                return col
            if value == "검증 결과":
                return col
    return None


def extract_status_map(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    status_col = _detect_status_column(ws)
    if status_col is None:
        raise RuntimeError(f"status column not found: {path}")

    item_col = 4
    status_map = {}
    for row in range(1, ws.max_row + 1):
        item = str(ws.cell(row=row, column=item_col).value or "").strip()
        status = str(ws.cell(row=row, column=status_col).value or "").strip()
        if not item or status not in VALID_STATUS:
            continue
        status_map[item] = status

    return status_map, status_col


def _classify_mismatch_reason(reference_status, generated_status):
    if reference_status == generated_status:
        return ""
    if generated_status is None:
        return "generated_missing_item"
    if generated_status == "N/A" and reference_status in {"NG", "OK"}:
        return "na_policy_gap"
    if reference_status == "NG" and generated_status in {"OK", "N/A"}:
        return "rule_missed_or_mapping_gap"
    if reference_status == "OK" and generated_status == "NG":
        return "potential_false_positive"
    return "status_policy_difference"


def compare(reference_path, generated_path):
    ref_map, ref_col = extract_status_map(reference_path)
    gen_map, gen_col = extract_status_map(generated_path)

    compare_keys = sorted([k for k, v in ref_map.items() if v in VALID_STATUS])
    rows = []
    mismatch = 0
    for item in compare_keys:
        ref_status = ref_map.get(item)
        gen_status = gen_map.get(item)
        is_match = ref_status == gen_status
        if not is_match:
            mismatch += 1
        rows.append(
            {
                "item": item,
                "reference": ref_status,
                "generated": gen_status,
                "match": is_match,
                "mismatch_reason": _classify_mismatch_reason(ref_status, gen_status),
            }
        )

    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "reference_path": os.path.abspath(reference_path),
        "generated_path": os.path.abspath(generated_path),
        "reference_status_column": ref_col,
        "generated_status_column": gen_col,
        "compare_item_count": len(compare_keys),
        "match_count": len(compare_keys) - mismatch,
        "mismatch_count": mismatch,
        "rows": rows,
    }


def main():
    parser = argparse.ArgumentParser(description="Compare reference review sheet vs generated review sheet.")
    parser.add_argument("--reference", required=True, help="Reference Excel file path")
    parser.add_argument("--generated", required=True, help="Generated Excel file path")
    parser.add_argument("--output", default="", help="Optional output JSON path")
    args = parser.parse_args()

    result = compare(args.reference, args.generated)
    text = json.dumps(result, ensure_ascii=False, indent=2)

    if args.output:
        os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"[+] Compare report written: {os.path.abspath(args.output)}")
    else:
        print(text)


if __name__ == "__main__":
    main()
